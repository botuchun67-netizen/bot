import asyncio
import logging
import os
import re
import sqlite3
import sys
from datetime import datetime, timedelta, timezone
from functools import wraps
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from telegram import InlineKeyboardButton as Btn
from telegram import InlineKeyboardMarkup as Kb
from telegram import Update
from telegram.error import BadRequest, Conflict, NetworkError, RetryAfter
from telegram.ext import (Application, ApplicationBuilder, CallbackQueryHandler,
                          CommandHandler, ContextTypes, MessageHandler,
                          PersistenceInput, PicklePersistence, filters)

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger("vote_bot")

# ===================== SOZLAMALAR =====================
TOKEN = os.getenv("BOT_TOKEN", "BOT_TOKEN_SHU_YERGA")
ADMIN_ID = 8355611778
TZ = timezone(timedelta(hours=5), "Toshkent")  # O'zbekiston: UTC+5, yozgi vaqt yo'q

# Railway Volume ulangan bo'lsa, ma'lumotlar shu papkada saqlanadi
DATA_DIR = os.getenv("RAILWAY_VOLUME_MOUNT_PATH") or os.getenv("DATA_DIR", ".")
os.makedirs(DATA_DIR, exist_ok=True)
DB = os.path.join(DATA_DIR, "voting.db")  # yangi fayl: ovozlar 0 dan boshlanadi
STATE_FILE = os.path.join(DATA_DIR, "state.pickle")

# Har ovozdan keyin bot qayta ishga tushsinmi. Ovozlar bazaga darhol yoziladi,
# restart kerak emas va u botni sekinlashtiradi. Faqat kerak bo'lsa True qiling.
RESTART_AFTER_VOTE = False
# False: oddiy foydalanuvchilarga Excelda faqat nomzodlar natijasi (kim ovoz bergani yashirin)
PUBLIC_SHOWS_VOTERS = False

CATEGORIES = ["💻 IT", "📚 Ta'lim", "⚽ Sport", "🎬 Media",
              "🎭 Madaniyat", "🌿 Ekologiya", "⚖️ Ombudsman"]

# Nomzodlar (tartib CATEGORIES bilan bir xil). Baza bo'sh bo'lganda BIR MARTA qo'shiladi;
# keyin nomzodlarni admin paneldan qo'shish/o'chirish mumkin.
CANDIDATES = [
    # 💻 IT
    ["Elmurod Saliybayev", "Ruzimbetov Sanjar", "Saliybayeva Zarina"],
    # 📚 Ta'lim
    ["Chorshanbayev Sarvar", "Gaipnazarov Anvar", "Yusupova Charos"],
    # ⚽ Sport
    ["Ayitbayev O'ktam", "Kamarov Bunyod", "Mirzayev Shohrux", "Qamaraddinov Jahongir"],
    # 🎬 Media
    ["Jumaboyev Farxod", "Ramatullayev Oʻlmasbek", "Saliybayeva Shamsiya", "Xamrayev Umidjon"],
    # 🎭 Madaniyat
    ["Ulugʻbekova Sohiba", "Kenjayev Xurshid", "Toʻrabayeva Nodira"],
    # 🌿 Ekologiya
    ["Abdullayeva Shahnoza", "Jangirova Elnora"],
    # ⚖️ Ombudsman
    ["Mirzayev Mirzohid", "Sapayev Behruz"],
]
# ======================================================

assert len(CANDIDATES) == len(CATEGORIES), "CANDIDATES va CATEGORIES soni bir xil bo'lishi kerak"

admin = filters.User(ADMIN_ID)
BACK_MENU = [Btn("⬅️ Yo'nalishlar", callback_data="menu")]
PANEL_BACK = [Btn("⬅️ Admin panel", callback_data="a:panel")]
CANCEL = [Btn("❌ Bekor qilish", callback_data="a:panel")]
CAP_PUB = "📊 Natijalar (Excel)"


# ---------- Baza (bitta ulanish, WAL, autocommit) ----------

_con = None
_end = None  # tugash vaqti xotirada keshlangan


def open_db():
    global _con
    _con = sqlite3.connect(DB, check_same_thread=False, isolation_level=None)
    try:
        _con.execute("PRAGMA journal_mode=WAL")
        _con.execute("PRAGMA synchronous=NORMAL")
    except sqlite3.DatabaseError:
        log.warning("WAL yoqilmadi, oddiy rejimda davom etadi")


def ex(sql, args=()):
    return _con.execute(sql, args)


def one(sql, args=()):
    return _con.execute(sql, args).fetchone()


def allr(sql, args=()):
    return _con.execute(sql, args).fetchall()


def init_db():
    _con.executescript("""
    CREATE TABLE IF NOT EXISTS candidates(
        id INTEGER PRIMARY KEY AUTOINCREMENT, cat INTEGER NOT NULL, name TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS votes(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        voter_key TEXT UNIQUE NOT NULL,
        voter_name TEXT NOT NULL,
        cand_id INTEGER NOT NULL,
        tg_id INTEGER, tg_username TEXT, ts TEXT);
    CREATE TABLE IF NOT EXISTS users(user_id INTEGER PRIMARY KEY);
    CREATE TABLE IF NOT EXISTS settings(k TEXT PRIMARY KEY, v TEXT);
    CREATE INDEX IF NOT EXISTS idx_votes_cand ON votes(cand_id);
    CREATE INDEX IF NOT EXISTS idx_cand_cat ON candidates(cat);
    """)
    if one("SELECT COUNT(*) FROM candidates")[0] == 0:  # birinchi ishga tushish
        for cat, names in enumerate(CANDIDATES, 1):
            for name in names:
                ex("INSERT INTO candidates(cat, name) VALUES(?,?)", (cat, name))
        log.info("Nomzodlar bazaga qo'shildi: %d ta", sum(len(n) for n in CANDIDATES))


def load_end():
    global _end
    r = one("SELECT v FROM settings WHERE k='end'")
    _end = datetime.fromisoformat(r[0]) if r else None


def get_end():
    return _end


def save_end(dt):
    global _end
    ex("INSERT OR REPLACE INTO settings VALUES('end', ?)", (dt.isoformat(),))
    _end = dt


def is_ended():
    return _end is not None and datetime.now(TZ) >= _end


# ---------- Ism-familiya ----------

WORD = re.compile(r"[^\W\d_]+(?:['\-][^\W\d_]+)*")


def clean_name(raw):
    """(ko'rinadigan ism, noyob kalit) yoki None. Ism-familiya tartibi farq qilmaydi."""
    s = raw.strip()
    for ch in "ʻʼ’‘`´":
        s = s.replace(ch, "'")
    words = s.split()
    if not 2 <= len(words) <= 5 or len(s) > 60:
        return None
    if not all(len(w) >= 2 and WORD.fullmatch(w) for w in words):
        return None
    display = " ".join(w[:1].upper() + w[1:] for w in words)
    key = " ".join(sorted(w.casefold() for w in words))
    return display, key


# ---------- Natijalar ----------

def results_text():
    rows = allr("""
        SELECT c.cat, c.name, COUNT(v.id) n FROM candidates c
        LEFT JOIN votes v ON v.cand_id = c.id
        GROUP BY c.id ORDER BY c.cat, n DESC, c.name""")
    out = ["🏆 Ovoz berish natijalari"]
    for i, name in enumerate(CATEGORIES, 1):
        out.append(f"\n📌 {name}")
        items = [r for r in rows if r[0] == i]
        if not items:
            out.append("  — nomzod yo'q")
        for _, cname, n in items:
            out.append(f"  {cname} — {n} ovoz")
    return "\n".join(out)


def _style(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5597")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(width + 3, 60)
    ws.freeze_panes = "A2"


def make_excel(full):
    """Alohida ulanish bilan ishlaydi (thread ichida chaqiriladi).
    full=True: ovoz berganlar ismlari va ID lari bilan to'liq hisobot."""
    con = sqlite3.connect(DB)
    try:
        summary = con.execute("""
            SELECT c.id, c.cat, c.name, COUNT(v.id) FROM candidates c
            LEFT JOIN votes v ON v.cand_id = c.id GROUP BY c.id
            ORDER BY c.cat, COUNT(v.id) DESC, c.name""").fetchall()
        voters = con.execute("""
            SELECT v.cand_id, v.voter_name, v.ts, v.tg_id, v.tg_username, c.cat, c.name
            FROM votes v JOIN candidates c ON c.id = v.cand_id ORDER BY v.id""").fetchall()
    finally:
        con.close()

    by_cand = {}
    for cid, vname, *_ in voters:
        by_cand.setdefault(cid, []).append(vname)

    wb = Workbook()
    ws = wb.active
    ws.title = "Natijalar"
    ws.append(["Yo'nalish", "O'rin", "Nomzod", "Ovozlar soni"] + (["Ovoz berganlar"] if full else []))
    prev_cat, prev_n, rank, pos = None, None, 0, 0
    for cid, cat, name, n in summary:
        if cat != prev_cat:
            prev_cat, prev_n, pos = cat, None, 0
        pos += 1
        if n != prev_n:
            rank, prev_n = pos, n
        row = [CATEGORIES[cat - 1], rank, name, n]
        if full:
            row.append(", ".join(by_cand.get(cid, [])))
        ws.append(row)
    ws.append([])
    ws.append(["Jami ovozlar", "", "", len(voters)])
    _style(ws)
    if full:
        for r in ws.iter_rows(min_row=2, min_col=5, max_col=5):
            r[0].alignment = Alignment(wrap_text=True, vertical="top")
        w2 = wb.create_sheet("Ovoz berganlar")
        w2.append(["№", "Ism-familiya", "Yo'nalish", "Nomzod", "Vaqt", "Telegram ID", "Username"])
        for i, (_, vname, ts, tg_id, un, cat, cname) in enumerate(voters, 1):
            t = datetime.fromisoformat(ts).strftime("%d.%m.%Y %H:%M:%S") if ts else ""
            w2.append([i, vname, CATEGORIES[cat - 1], cname, t, tg_id, f"@{un}" if un else ""])
        _style(w2)

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


async def send_results(bot, chat_id, full):
    data = await asyncio.to_thread(make_excel, full)
    name = "natijalar_toliq.xlsx" if full else "natijalar.xlsx"
    cap = "📊 To'liq natijalar (ovoz berganlar bilan)" if full else CAP_PUB
    await bot.send_document(chat_id, data, filename=name, caption=cap)


# ---------- Yordamchilar ----------

def cat_rows(prefix):
    btns = [Btn(n, callback_data=f"{prefix}:{i}") for i, n in enumerate(CATEGORIES, 1)]
    return [btns[j:j + 2] for j in range(0, len(btns), 2)]


def main_kb(uid):
    rows = cat_rows("cat")
    rows.append([Btn("📊 Natijalar", callback_data="res")])
    if uid == ADMIN_ID:
        rows.append([Btn("⚙️ Admin panel", callback_data="a:panel")])
    return Kb(rows)


async def ans(q, text=None, alert=False):
    """Tugma bosilganini darhol tasdiqlaydi; eski so'rovlarda xato bermaydi."""
    try:
        await q.answer(text, show_alert=alert)
    except Exception:
        pass


async def show(q, text, kb=None):
    """Xabarni tahrirlaydi; tahrirlab bo'lmasa yangisini yuboradi."""
    try:
        await q.message.edit_text(text, reply_markup=kb)
    except BadRequest as e:
        if "not modified" in str(e).lower():
            return
        await q.message.reply_text(text, reply_markup=kb)


MAIN_TEXT = "Yo'nalishni tanlang.\n⚠️ Har bir ism-familiya faqat 1 marta ovoz bera oladi."


# ---------- Foydalanuvchi ----------

async def start(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    ctx.user_data.clear()
    uid = update.effective_user.id
    ex("INSERT OR IGNORE INTO users VALUES(?)", (uid,))
    text = "Ovoz berish yakunlangan." if is_ended() else "Salom! " + MAIN_TEXT
    await update.message.reply_text(text, reply_markup=main_kb(uid))


async def menu(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    await ans(q)
    ctx.user_data.clear()
    await show(q, MAIN_TEXT, main_kb(q.from_user.id))


async def show_cat(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    if is_ended():
        await ans(q, "Ovoz berish yakunlangan.", True)
        return
    await ans(q)
    ctx.user_data.clear()
    cat = int(q.data.split(":")[1])
    title = CATEGORIES[cat - 1]
    cands = allr("SELECT id, name FROM candidates WHERE cat=? ORDER BY id", (cat,))
    if not cands:
        await show(q, f"{title}\n\nBu yo'nalishda hozircha nomzod yo'q.", Kb([BACK_MENU]))
        return
    rows = [[Btn(name, callback_data=f"v:{cid}")] for cid, name in cands]
    rows.append(BACK_MENU)
    await show(q, f"📌 {title}\n\nOvoz bermoqchi bo'lgan nomzodni tanlang:", Kb(rows))


async def vote(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """1-qadam: nomzod tanlandi, ism-familiya so'raladi."""
    q = update.callback_query
    if is_ended():
        await ans(q, "Ovoz berish yakunlangan.", True)
        return
    cid = int(q.data.split(":")[1])
    cand = one("SELECT name, cat FROM candidates WHERE id=?", (cid,))
    if not cand:
        await ans(q, "Nomzod topilmadi.", True)
        return
    await ans(q)
    ctx.user_data.clear()
    ctx.user_data.update(vstate="name", vcand=cid)
    await show(q, f"🗳 Nomzod: {cand[0]}\n\n✍️ Ism va familiyangizni yozing (masalan: Jasur Karimov).\n"
                  "⚠️ Har bir ism-familiya faqat 1 marta ovoz bera oladi, keyin o'zgartirib bo'lmaydi.",
               Kb([[Btn("❌ Bekor qilish", callback_data=f"cat:{cand[1]}")]]))


async def vote_name(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """2-qadam: ism-familiya yozildi, tasdiqlash so'raladi."""
    m, ud = update.message, ctx.user_data
    cancel = Kb([[Btn("❌ Bekor qilish", callback_data="menu")]])
    if is_ended():
        ud.clear()
        await m.reply_text("Ovoz berish yakunlangan.")
        return
    parsed = clean_name(m.text)
    if not parsed:
        await m.reply_text("❗ Ism va familiyangizni to'liq, harflar bilan yozing.\nMasalan: Jasur Karimov",
                           reply_markup=cancel)
        return
    name, key = parsed
    taken = one("SELECT 1 FROM votes WHERE voter_key=?", (key,))
    cand = one("SELECT name, cat FROM candidates WHERE id=?", (ud.get("vcand"),))
    if taken:
        await m.reply_text("⛔ Bu ism-familiya bilan allaqachon ovoz berilgan.\n"
                           "Boshqa ism-familiya yozing yoki bekor qiling.", reply_markup=cancel)
        return
    if not cand:
        ud.clear()
        await m.reply_text("Nomzod topilmadi.", reply_markup=Kb([BACK_MENU]))
        return
    ud.update(vstate="confirm", vname=name, vkey=key)
    kb = Kb([[Btn("✅ Tasdiqlash", callback_data="vc:yes"), Btn("❌ Bekor qilish", callback_data="menu")]])
    await m.reply_text(f"👤 Ovoz beruvchi: {name}\n📌 {CATEGORIES[cand[1] - 1]}\n🗳 Nomzod: {cand[0]}\n\n"
                       "⚠️ Ovoz qabul qilingach, o'zgartirib bo'lmaydi.", reply_markup=kb)


async def vote_confirm(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    """3-qadam: ovozni bazaga yozish. Holat darhol tozalanadi, ikki marta bosish ovozni takrorlamaydi."""
    q, ud = update.callback_query, ctx.user_data
    if is_ended():
        await ans(q, "Ovoz berish yakunlangan.", True)
        return
    if ud.get("vstate") != "confirm":
        await ans(q, "Bu ovoz allaqachon hisoblangan yoki sessiya tugagan. /start bosing.", True)
        return
    cid, name, key = ud["vcand"], ud["vname"], ud["vkey"]
    ud.clear()
    if not one("SELECT 1 FROM candidates WHERE id=?", (cid,)):
        await ans(q, "Nomzod topilmadi.", True)
        return
    try:
        ex("""INSERT INTO votes(voter_key, voter_name, cand_id, tg_id, tg_username, ts)
              VALUES(?,?,?,?,?,?)""",
           (key, name, cid, q.from_user.id, q.from_user.username, datetime.now(TZ).isoformat()))
    except sqlite3.IntegrityError:
        await ans(q, "Bu ism-familiya bilan allaqachon ovoz berilgan.", True)
        return
    await ans(q, "✅ Ovozingiz qabul qilindi!", True)
    await show(q, f"✅ {name}, ovozingiz qabul qilindi.\nNatijalar ovoz berish tugagach e'lon qilinadi.",
               Kb([BACK_MENU]))
    if RESTART_AFTER_VOTE:
        schedule_restart(ctx.application)


async def results(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    is_admin = q.from_user.id == ADMIN_ID
    if not (is_ended() or is_admin):
        await ans(q, "Natijalar ovoz berish tugagach e'lon qilinadi.", True)
        return
    await ans(q)
    await show(q, results_text(), Kb([BACK_MENU]))
    await send_results(ctx.bot, q.message.chat_id, is_admin or PUBLIC_SHOWS_VOTERS)


# ---------- Admin ----------

def admin_only(fn):
    @wraps(fn)
    async def wrapper(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
        if update.effective_user.id != ADMIN_ID:
            if update.callback_query:
                await ans(update.callback_query, "Ruxsat yo'q", True)
            return
        return await fn(update, ctx)
    return wrapper


def panel_text():
    users = one("SELECT COUNT(*) FROM users")[0]
    votes = one("SELECT COUNT(*) FROM votes")[0]
    cands = one("SELECT COUNT(*) FROM candidates")[0]
    end = get_end()
    end_s = f"{end:%d.%m.%Y %H:%M}" if end else "belgilanmagan"
    return (f"⚙️ Admin panel\n\n👥 Foydalanuvchilar: {users}\n🗳 Ovozlar: {votes}\n"
            f"👤 Nomzodlar: {cands}\n⏰ Tugash: {end_s}")


def panel_kb():
    return Kb([
        [Btn("➕ Nomzod qo'shish", callback_data="a:add"), Btn("📋 Nomzodlar", callback_data="a:list")],
        [Btn("⏰ Tugash vaqti", callback_data="a:end"), Btn("📊 Natijalar", callback_data="a:res")],
        [Btn("🧹 Ovozlarni nolga tushirish", callback_data="a:reset")],
        [Btn("⬅️ Asosiy menyu", callback_data="menu")],
    ])


@admin_only
async def admin_cmd(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    ctx.user_data.clear()
    await update.message.reply_text(panel_text(), reply_markup=panel_kb())


@admin_only
async def admin_cb(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    await ans(q)
    p = q.data.split(":")
    act = p[1]

    if act == "panel":
        ctx.user_data.clear()
        await show(q, panel_text(), panel_kb())

    elif act == "add":
        await show(q, "Qaysi yo'nalishga nomzod qo'shamiz?", Kb(cat_rows("a:addcat") + [PANEL_BACK]))

    elif act == "addcat":
        n = int(p[2])
        ctx.user_data.clear()
        ctx.user_data.update(state="add_name", cat=n)
        await show(q, f"{CATEGORIES[n - 1]}\n\n✍️ Nomzodning ism-familiyasini yozing:", Kb([CANCEL]))

    elif act == "list":
        await show(q, "Qaysi yo'nalish nomzodlari?", Kb(cat_rows("a:listcat") + [PANEL_BACK]))

    elif act == "listcat":
        n = int(p[2])
        rows = allr("""SELECT c.id, c.name, COUNT(v.id) FROM candidates c
                       LEFT JOIN votes v ON v.cand_id = c.id WHERE c.cat=? GROUP BY c.id ORDER BY c.id""", (n,))
        if not rows:
            await show(q, f"{CATEGORIES[n - 1]}: nomzod yo'q.", Kb([PANEL_BACK]))
            return
        kb = [[Btn(f"🗑 {name} ({cnt})", callback_data=f"a:delask:{cid}")] for cid, name, cnt in rows]
        await show(q, f"{CATEGORIES[n - 1]}\n\nO'chirish uchun nomzodni bosing (qavsda ovozlar soni):",
                   Kb(kb + [PANEL_BACK]))

    elif act == "delask":
        cid = int(p[2])
        r = one("SELECT name, cat FROM candidates WHERE id=?", (cid,))
        if not r:
            await show(q, "Nomzod topilmadi.", Kb([PANEL_BACK]))
            return
        kb = Kb([[Btn("✅ Ha, o'chirish", callback_data=f"a:del:{cid}"),
                  Btn("❌ Yo'q", callback_data=f"a:listcat:{r[1]}")]])
        await show(q, f"🗑 {r[0]} o'chirilsinmi?\nUning ovozlari ham o'chadi.", kb)

    elif act == "del":
        cid = int(p[2])
        r = one("SELECT cat FROM candidates WHERE id=?", (cid,))
        if r:
            ex("DELETE FROM votes WHERE cand_id=?", (cid,))
            ex("DELETE FROM candidates WHERE id=?", (cid,))
        back = [Btn("⬅️ Ro'yxat", callback_data=f"a:listcat:{r[0]}")] if r else []
        await show(q, "🗑 O'chirildi", Kb([back, PANEL_BACK] if back else [PANEL_BACK]))

    elif act == "end":
        kb = Kb([
            [Btn("+1 kun", callback_data="a:endin:1"), Btn("+3 kun", callback_data="a:endin:3"),
             Btn("+7 kun", callback_data="a:endin:7")],
            [Btn("✍️ Sana va vaqtni yozish", callback_data="a:endman")],
            [Btn("🏁 Hozir yakunlash", callback_data="a:endnow")],
            PANEL_BACK,
        ])
        await show(q, "⏰ Ovoz berish qachon tugasin? (hozirdan boshlab)", kb)

    elif act == "endin":
        dt = datetime.now(TZ) + timedelta(days=int(p[2]))
        save_end(dt)
        schedule_finish(ctx.application)
        await show(q, f"✅ Tugash vaqti: {dt:%d.%m.%Y %H:%M}", Kb([PANEL_BACK]))

    elif act == "endman":
        ctx.user_data.clear()
        ctx.user_data["state"] = "set_end"
        await show(q, "Sana va vaqtni yozing (Toshkent vaqti):\n2026-10-05 18:00", Kb([CANCEL]))

    elif act == "endnow":
        kb = Kb([[Btn("✅ Ha, yakunlash", callback_data="a:endyes"),
                  Btn("❌ Yo'q", callback_data="a:panel")]])
        await show(q, "⚠️ Ovoz berish hoziroq tugaydi va natijalar hammaga yuboriladi. Tasdiqlaysizmi?", kb)

    elif act == "endyes":
        save_end(datetime.now(TZ))
        schedule_finish(ctx.application)
        await show(q, "🏁 Ovoz berish yakunlandi. Natijalar yuborilmoqda.", Kb([PANEL_BACK]))

    elif act == "res":
        await show(q, results_text(), Kb([PANEL_BACK]))
        await send_results(ctx.bot, q.message.chat_id, True)

    elif act == "reset":
        n = one("SELECT COUNT(*) FROM votes")[0]
        kb = Kb([[Btn("✅ Ha, nolga tushirish", callback_data="a:resetyes"),
                  Btn("❌ Yo'q", callback_data="a:panel")]])
        await show(q, f"⚠️ Barcha {n} ta ovoz o'chiriladi va hisob 0 dan boshlanadi.\n"
                      "Nomzodlar va tugash vaqti saqlanadi. Tasdiqlaysizmi?", kb)

    elif act == "resetyes":
        ex("DELETE FROM votes")
        await show(q, "✅ Barcha ovozlar o'chirildi. Hisob 0 dan boshlandi.", Kb([PANEL_BACK]))


async def admin_text(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    m = update.message
    state = ctx.user_data.get("state")
    if state == "add_name":
        name = " ".join(m.text.split())
        if not name or len(name) > 60:
            await m.reply_text("Ism 1–60 belgi bo'lsin. Qayta yozing:")
            return
        n = ctx.user_data["cat"]
        ex("INSERT INTO candidates(cat, name) VALUES(?,?)", (n, name))
        ctx.user_data.clear()
        kb = Kb([[Btn("➕ Yana qo'shish", callback_data=f"a:addcat:{n}")], PANEL_BACK])
        await m.reply_text(f"✅ Qo'shildi: {name} → {CATEGORIES[n - 1]}", reply_markup=kb)
    elif state == "set_end":
        try:
            dt = datetime.strptime(m.text.strip(), "%Y-%m-%d %H:%M").replace(tzinfo=TZ)
        except ValueError:
            await m.reply_text("Format noto'g'ri. Masalan: 2026-10-05 18:00")
            return
        save_end(dt)
        schedule_finish(ctx.application)
        ctx.user_data.clear()
        await m.reply_text(f"✅ Tugash vaqti: {dt:%d.%m.%Y %H:%M}", reply_markup=Kb([PANEL_BACK]))


async def on_text(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    ud = ctx.user_data
    if update.effective_user.id == ADMIN_ID and ud.get("state") in ("add_name", "set_end"):
        await admin_text(update, ctx)
    elif ud.get("vstate") == "name":
        await vote_name(update, ctx)
    else:
        await update.message.reply_text("Boshlash uchun /start bosing.")


# ---------- Yakunlash, restart, xatolar ----------

def _secs(ra):
    return ra.total_seconds() if hasattr(ra, "total_seconds") else float(ra)


async def _notify(bot, uid, text, doc, box):
    for _ in range(2):
        try:
            await bot.send_message(uid, text)
            if box["fid"]:
                await bot.send_document(uid, box["fid"], caption=CAP_PUB)
            else:
                msg = await bot.send_document(uid, doc, filename="natijalar.xlsx", caption=CAP_PUB)
                box["fid"] = msg.document.file_id
            return
        except RetryAfter as e:
            await asyncio.sleep(_secs(e.retry_after) + 1)
        except Exception:
            return  # bloklagan yoki o'chirilgan foydalanuvchi


async def finish(ctx: ContextTypes.DEFAULT_TYPE):
    text = results_text()
    public = await asyncio.to_thread(make_excel, PUBLIC_SHOWS_VOTERS)
    users = [r[0] for r in allr("SELECT user_id FROM users") if r[0] != ADMIN_ID]
    box = {"fid": None}
    if users:
        await _notify(ctx.bot, users[0], text, public, box)  # birinchisi faylni yuklab file_id beradi
        for i in range(1, len(users), 10):
            await asyncio.gather(*(_notify(ctx.bot, u, text, public, box) for u in users[i:i + 10]))
            await asyncio.sleep(1)  # Telegram limitidan oshmaslik uchun
    try:
        await ctx.bot.send_message(ADMIN_ID, text)
        await send_results(ctx.bot, ADMIN_ID, True)
    except Exception:
        log.exception("Adminga natijani yuborib bo'lmadi")


def schedule_finish(app: Application):
    for job in app.job_queue.get_jobs_by_name("finish"):
        job.schedule_removal()
    end = get_end()
    if end:
        delay = max((end - datetime.now(TZ)).total_seconds(), 1)
        app.job_queue.run_once(finish, delay, name="finish")


def schedule_restart(app: Application):
    if not app.job_queue.get_jobs_by_name("restart"):
        app.job_queue.run_once(restart, 3, name="restart")


async def restart(ctx: ContextTypes.DEFAULT_TYPE):
    script = sys.argv[0] if sys.argv else ""
    if not (script and os.path.isfile(script)):
        log.warning("Avto-restart o'tkazib yuborildi: skript yo'li aniqlanmadi.")
        return
    try:
        await ctx.application.update_persistence()
        log.info("Bot qayta ishga tushmoqda...")
        os.execv(sys.executable, [sys.executable] + sys.argv)
    except Exception:
        log.exception("Avto-restart ishlamadi, bot davom etadi.")


async def on_error(update: object, ctx: ContextTypes.DEFAULT_TYPE):
    err = ctx.error
    if isinstance(err, Conflict):
        log.error("⚠️ Bot BOSHQA JOYDA ham ishlayapti (Conflict). Eski nusxani to'xtating "
                  "yoki BotFather'da /revoke qilib tokenni almashtiring.")
    elif isinstance(err, NetworkError):
        log.warning("Tarmoq xatosi: %s", err)
    else:
        log.error("Kutilmagan xato", exc_info=err)


async def post_init(app: Application):
    schedule_finish(app)


def main():
    if TOKEN.startswith("BOT_TOKEN_SHU"):
        sys.exit("BOT_TOKEN o'rnatilmagan: Railway → Variables ga BOT_TOKEN qo'shing.")
    if (os.getenv("RAILWAY_ENVIRONMENT_NAME") or os.getenv("RAILWAY_ENVIRONMENT")) \
            and not os.getenv("RAILWAY_VOLUME_MOUNT_PATH"):
        log.warning("Volume ulanmagan: qayta deploydan keyin ovozlar yo'qoladi!")
    open_db()
    init_db()
    load_end()

    builder = (ApplicationBuilder().token(TOKEN)
               .concurrent_updates(True)      # foydalanuvchilar bir-birini kutmaydi
               .connection_pool_size(64)
               .post_init(post_init))
    if RESTART_AFTER_VOTE:
        builder = builder.persistence(PicklePersistence(
            STATE_FILE, update_interval=5,
            store_data=PersistenceInput(bot_data=False, chat_data=False, callback_data=False)))
    app = builder.build()

    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("admin", admin_cmd, filters=admin))
    app.add_handler(CallbackQueryHandler(menu, pattern=r"^menu$"))
    app.add_handler(CallbackQueryHandler(show_cat, pattern=r"^cat:\d+$"))
    app.add_handler(CallbackQueryHandler(vote, pattern=r"^v:\d+$"))
    app.add_handler(CallbackQueryHandler(vote_confirm, pattern=r"^vc:yes$"))
    app.add_handler(CallbackQueryHandler(results, pattern=r"^res$"))
    app.add_handler(CallbackQueryHandler(admin_cb, pattern=r"^a:"))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, on_text))
    app.add_error_handler(on_error)

    log.info("Bot ishga tushdi")
    app.run_polling(allowed_updates=["message", "callback_query"], drop_pending_updates=True)


if __name__ == "__main__":
    main()
